from __future__ import annotations

import sqlite3
from datetime import date, datetime
from pathlib import Path

import openpyxl

from config.settings import INTRADAY_CHECKPOINTS, INTRADAY_REPORTS_DIR

HOUR_BLOCK_COLUMNS = ["Predicted", "Actual", "Error %", "Direction", "Confidence %"]


def _fetch_today_predictions(conn: sqlite3.Connection, run_date: date) -> list[sqlite3.Row]:
    return conn.execute(
        "SELECT p.*, a.actual_price, a.pct_error, a.direction_correct "
        "FROM intraday_predictions p LEFT JOIN intraday_accuracy_evaluations a ON a.prediction_id = p.id "
        "WHERE date(p.prediction_timestamp) = ? ORDER BY p.symbol, p.prediction_timestamp",
        (run_date.isoformat(),),
    ).fetchall()


def _fetch_open_prices(conn: sqlite3.Connection, symbols: list[str], run_date: date) -> dict[str, float]:
    prices: dict[str, float] = {}
    for symbol in symbols:
        row = conn.execute(
            "SELECT raw_current_price FROM intraday_predictions "
            "WHERE symbol = ? AND date(prediction_timestamp) = ? "
            "AND substr(prediction_timestamp, 12, 5) = '09:15' LIMIT 1",
            (symbol, run_date.isoformat()),
        ).fetchone()
        if row:
            prices[symbol] = row["raw_current_price"]
    return prices


def generate_intraday_report(
    conn: sqlite3.Connection, run_date: date, symbols: list[str], run_timestamp: datetime,
    market_status: str, processed: int, skipped: int, output_dir: Path = INTRADAY_REPORTS_DIR,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{run_date.isoformat()}.xlsx"
    tmp_path = output_dir / f".{run_date.isoformat()}.xlsx.tmp"

    rows = _fetch_today_predictions(conn, run_date)
    by_symbol: dict[str, dict[str, sqlite3.Row]] = {}
    for row in rows:
        checkpoint = row["evaluation_timestamp"][11:16]
        by_symbol.setdefault(row["symbol"], {})[checkpoint] = row

    open_prices = _fetch_open_prices(conn, symbols, run_date)

    evaluated = [r for r in rows if r["actual_price"] is not None]
    confidences = [r["confidence"] for r in rows]
    avg_confidence = sum(confidences) / len(confidences) if confidences else None
    direction_hits = [r["direction_correct"] for r in evaluated]
    direction_accuracy = (sum(direction_hits) / len(direction_hits) * 100) if direction_hits else None
    errors = [r["pct_error"] for r in evaluated]
    avg_error = sum(errors) / len(errors) if errors else None
    total_tokens = sum(r["input_tokens"] + r["output_tokens"] for r in rows)

    best = min(evaluated, key=lambda r: r["pct_error"], default=None)
    worst = max(evaluated, key=lambda r: r["pct_error"], default=None)

    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Intraday"

    ws.append([
        f"Run: {run_timestamp.isoformat()}", f"Market: {market_status}",
        f"Processed: {processed}", f"Skipped: {skipped}",
    ])
    ws.append([
        f"Avg Confidence: {avg_confidence:.0f}%" if avg_confidence is not None else "Avg Confidence: n/a",
        f"Direction Acc so far: {direction_accuracy:.0f}%" if direction_accuracy is not None else "Direction Acc so far: n/a",
        f"Avg Error: {avg_error:.2f}%" if avg_error is not None else "Avg Error: n/a",
        f"Tokens used today: {total_tokens}",
    ])
    ws.append([
        f"Best: {best['symbol']} {best['pct_error']:.2f}% err" if best else "Best: n/a",
        f"Worst: {worst['symbol']} {worst['pct_error']:.2f}% err" if worst else "Worst: n/a",
    ])
    ws.append([
        "Indicators: SMA(5/10/20/50h) EMA(9/21h) RSI(14h) MACD(12/26/9h) BB(20h,2sd) | "
        "Checkpoints: 09:15-15:15 hourly | Provider: yfinance",
    ])
    ws.append([])

    hour_checkpoints = INTRADAY_CHECKPOINTS[1:]
    header_top = ["Stock", "09:15"]
    header_sub = ["", "Actual"]
    for cp in hour_checkpoints:
        header_top.extend([cp, "", "", "", ""])
        header_sub.extend(HOUR_BLOCK_COLUMNS)
    ws.append(header_top)
    ws.append(header_sub)

    header_row = ws.max_row - 1
    col = 3
    for _ in hour_checkpoints:
        ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row, end_column=col + 4)
        col += 5

    for symbol in symbols:
        checkpoints = by_symbol.get(symbol, {})
        row_cells = [symbol, open_prices.get(symbol, "")]
        for cp in hour_checkpoints:
            record = checkpoints.get(cp)
            if record is None:
                row_cells.extend(["", "", "", "", ""])
                continue
            if record["direction_correct"] is None:
                direction_mark = ""
            else:
                direction_mark = "Yes" if record["direction_correct"] else "No"
            row_cells.extend([
                record["predicted_price"],
                record["actual_price"] if record["actual_price"] is not None else "",
                f"{record['pct_error']:.2f}%" if record["pct_error"] is not None else "",
                direction_mark,
                f"{record['confidence']:.0f}%",
            ])
        ws.append(row_cells)

    ws.freeze_panes = f"C{header_row + 2}"

    workbook.save(tmp_path)
    tmp_path.replace(output_path)
    return output_path
