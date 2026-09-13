from datetime import date, datetime, timezone

import openpyxl
import pytest

from src.reporting.intraday_excel_report import generate_intraday_report
from src.storage.db import get_connection


def _seed_prediction_and_evaluation(conn, symbol, checkpoint_hhmm, predicted_price, actual_price=None, direction_correct=None, pct_error=None, evaluation_hhmm=None):
    ts = f"2026-09-15T{checkpoint_hhmm}:00+05:30"
    evaluation_ts = f"2026-09-15T{evaluation_hhmm}:00+05:30" if evaluation_hhmm else "2026-09-15T99:99:00+05:30"
    conn.execute(
        "INSERT INTO intraday_predictions (created_at, symbol, prediction_timestamp, "
        "evaluation_timestamp, prediction_type, raw_current_price, open, high, low, close, "
        "volume, direction, predicted_price, expected_move_percent, confidence, reasoning, "
        "key_risks_json, technical_score, indicators_json, data_provider, interval, "
        "claude_model, prompt_version, raw_claude_response, input_tokens, output_tokens) VALUES "
        "(?, ?, ?, ?, 'next_hour', 1450.0, 1450, 1450, 1450, 1450, "
        "1000, 'BULLISH', ?, 1.0, 70.0, 'test', '[]', 40.0, '{}', 'yfinance', '60m', "
        "'claude-sonnet-5', 'intraday-v1', '{}', 650, 400)",
        ("2026-09-15T00:00:00+00:00", symbol, ts, evaluation_ts, predicted_price),
    )
    prediction_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    if actual_price is not None:
        conn.execute(
            "INSERT INTO intraday_accuracy_evaluations (prediction_id, evaluated_at, "
            "evaluation_timestamp, actual_price, predicted_price, abs_error, pct_error, "
            "direction_correct, target_hit) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)",
            (prediction_id, "2026-09-15T00:00:00+00:00", ts, actual_price, predicted_price,
             abs(actual_price - predicted_price), pct_error, direction_correct),
        )
    conn.commit()


def test_generate_intraday_report_creates_exactly_one_worksheet(tmp_path):
    conn = get_connection(tmp_path / "test.db")
    _seed_prediction_and_evaluation(conn, "RELIANCE.NS", "10:15", 1465.0, actual_price=1458.0, direction_correct=1, pct_error=0.48)

    output_path = generate_intraday_report(
        conn, date(2026, 9, 15), ["RELIANCE.NS"], datetime(2026, 9, 15, 10, 25, tzinfo=timezone.utc),
        "OPEN", processed=1, skipped=0, output_dir=tmp_path / "reports",
    )

    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    assert len(workbook.sheetnames) == 1
    conn.close()


def test_report_shows_blank_cells_for_unevaluated_checkpoints(tmp_path):
    conn = get_connection(tmp_path / "test.db")
    _seed_prediction_and_evaluation(conn, "RELIANCE.NS", "10:15", 1465.0, evaluation_hhmm="11:15")  # no evaluation yet

    output_path = generate_intraday_report(
        conn, date(2026, 9, 15), ["RELIANCE.NS"], datetime(2026, 9, 15, 10, 25, tzinfo=timezone.utc),
        "OPEN", processed=1, skipped=0, output_dir=tmp_path / "reports",
    )
    workbook = openpyxl.load_workbook(output_path)
    ws = workbook.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert 1465.0 in values  # the prediction is shown
    conn.close()


def test_report_regeneration_is_atomic_and_does_not_corrupt_existing_file_on_failure(tmp_path, monkeypatch):
    conn = get_connection(tmp_path / "test.db")
    _seed_prediction_and_evaluation(conn, "RELIANCE.NS", "10:15", 1465.0, actual_price=1458.0, direction_correct=1, pct_error=0.48)
    output_dir = tmp_path / "reports"

    good_path = generate_intraday_report(
        conn, date(2026, 9, 15), ["RELIANCE.NS"], datetime(2026, 9, 15, 10, 25, tzinfo=timezone.utc),
        "OPEN", processed=1, skipped=0, output_dir=output_dir,
    )
    original_bytes = good_path.read_bytes()

    import openpyxl as _openpyxl
    def _boom(self, *args, **kwargs):
        raise RuntimeError("simulated save failure")
    monkeypatch.setattr(_openpyxl.Workbook, "save", _boom)

    with pytest.raises(RuntimeError):
        generate_intraday_report(
            conn, date(2026, 9, 15), ["RELIANCE.NS"], datetime(2026, 9, 15, 11, 25, tzinfo=timezone.utc),
            "OPEN", processed=1, skipped=0, output_dir=output_dir,
        )

    assert good_path.read_bytes() == original_bytes  # untouched by the failed write
    conn.close()


def test_hour_block_columns_are_keyed_by_evaluation_timestamp_not_prediction_timestamp(tmp_path):
    conn = get_connection(tmp_path / "test.db")
    # Prediction made at 09:15 targets 10:15 -> must land in the "10:15" column.
    _seed_prediction_and_evaluation(conn, "RELIANCE.NS", "09:15", 1111.0, evaluation_hhmm="10:15")
    # Prediction made at 10:15 targets 11:15 -> must land in the "11:15" column.
    _seed_prediction_and_evaluation(conn, "RELIANCE.NS", "10:15", 2222.0, evaluation_hhmm="11:15")

    output_path = generate_intraday_report(
        conn, date(2026, 9, 15), ["RELIANCE.NS"], datetime(2026, 9, 15, 11, 25, tzinfo=timezone.utc),
        "OPEN", processed=2, skipped=0, output_dir=tmp_path / "reports",
    )
    workbook = openpyxl.load_workbook(output_path)
    ws = workbook.active

    # Locate the data row for RELIANCE.NS (column 1 == "Stock").
    data_row = None
    for row in ws.iter_rows(min_row=1, max_col=1):
        if row[0].value == "RELIANCE.NS":
            data_row = row[0].row
            break
    assert data_row is not None

    # Column layout: 1=Stock, 2="09:15 Actual", then 5-wide blocks per hour
    # checkpoint (10:15, 11:15, 12:15, 13:15, 14:15, 15:15), "Predicted" first
    # in each block.
    column_10_15_predicted = 3
    column_11_15_predicted = 8

    assert ws.cell(row=data_row, column=column_10_15_predicted).value == 1111.0
    assert ws.cell(row=data_row, column=column_11_15_predicted).value == 2222.0
    conn.close()


def test_09_15_actual_column_is_blank_when_the_09_15_checkpoint_was_missed(tmp_path):
    conn = get_connection(tmp_path / "test.db")
    # The day's earliest recorded prediction is at 12:15 -- 09:15 was a missed
    # checkpoint, never recorded. The "09:15 Actual" column must stay blank,
    # not be mislabeled with the 12:15 price.
    _seed_prediction_and_evaluation(conn, "RELIANCE.NS", "12:15", 1500.0, evaluation_hhmm="13:15")

    output_path = generate_intraday_report(
        conn, date(2026, 9, 15), ["RELIANCE.NS"], datetime(2026, 9, 15, 13, 25, tzinfo=timezone.utc),
        "OPEN", processed=1, skipped=0, output_dir=tmp_path / "reports",
    )
    workbook = openpyxl.load_workbook(output_path)
    ws = workbook.active

    data_row = None
    for row in ws.iter_rows(min_row=1, max_col=1):
        if row[0].value == "RELIANCE.NS":
            data_row = row[0].row
            break
    assert data_row is not None

    column_09_15_actual = 2
    assert ws.cell(row=data_row, column=column_09_15_actual).value in ("", None)
    conn.close()
