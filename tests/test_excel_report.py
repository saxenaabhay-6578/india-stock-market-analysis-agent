from datetime import date

import openpyxl

from src.reporting.excel_report import generate_report
from src.storage.db import get_connection

EXPECTED_SHEETS = {
    "Dashboard", "Today's Predictions", "Prediction History", "Accuracy",
    "Stock Performance", "Historical Prices", "News-Sentiment", "Configuration",
}


def _seed_minimal_data(conn):
    conn.execute(
        "INSERT INTO predictions (created_at, prediction_date, symbol, horizon_days, "
        "target_evaluation_date, current_price, direction, target_price, range_low, "
        "range_high, expected_move_percent, risk_level, confidence, reasoning, "
        "key_risks_json, technical_score, indicators_json, data_timestamp, "
        "data_provider, claude_model, prompt_version, raw_claude_response) VALUES "
        "('2026-09-14T00:00:00', '2026-09-14', 'TCS.NS', 1, '2026-09-15', 1000.0, "
        "'BULLISH', 1010.0, 1005.0, 1015.0, 1.0, 990.0, 70.0, 'test', '[]', 50.0, "
        "'{}', '2026-09-14T00:00:00', 'yfinance', 'claude-sonnet-5', 'v1', '{}')"
    )
    conn.execute(
        "INSERT INTO price_bars (symbol, date, open, high, low, close, adj_close, volume, source, fetched_at) "
        "VALUES ('TCS.NS', '2026-09-14', 1000, 1010, 995, 1005, 1005, 1000, 'yfinance', '2026-09-14T00:00:00')"
    )
    conn.commit()


def test_generate_report_creates_all_expected_sheets(tmp_path):
    conn = get_connection(tmp_path / "test.db")
    _seed_minimal_data(conn)

    output_path = generate_report(
        conn, date(2026, 9, 14), ["TCS.NS"], predictions_made=1, predictions_skipped=0,
        output_dir=tmp_path / "reports",
    )

    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    assert set(workbook.sheetnames) == EXPECTED_SHEETS
    dashboard = workbook["Dashboard"]
    assert dashboard.cell(row=1, column=1).value == "run_date"
    conn.close()
